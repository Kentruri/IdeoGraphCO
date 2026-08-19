"""Ingesta del gold anotado: Excel → α de Krippendorff → consenso → JSONL humano.

Este script cierra el hueco crítico del pipeline: antes de él, el "test gold"
se evaluaba contra las etiquetas SILVER del LLM porque la anotación humana
nunca volvía del Excel al pipeline.

Flujo (metodología del anteproyecto — Calibración Humana):
1. Lee los libros anotados (uno por anotador, hoja "Articulos").
2. Sobre el bloque de SOLAPE calcula Krippendorff α:
   - nominal sobre `clase_dominante` (el número que exige el anteproyecto, ≥ 0.8)
   - interval sobre cada eje 1-5 (diagnóstico del codebook)
3. Las discrepancias en la dominante van a un CSV para la sesión de consenso;
   se resuelven pasando `--consensus` con la columna `label_final` llena.
4. Escribe `annotation/gold_set_<v>_labeled.jsonl` con label/label_idx +
   label_source="human" — lo que consume scripts/prepare_splits.py.
5. (Opcional) `--audit-silver`: α y accuracy del juez LLM contra el consenso
   humano sobre los mismos artículos (auditoría del Silver Set, P1.3).

Uso:
    python scripts/ingest_gold.py --books annotation/gold_set_v2_kevin.xlsx \\
        annotation/gold_set_v2_juan.xlsx
    # tras la sesión de consenso:
    python scripts/ingest_gold.py --books ... --consensus annotation/gold_v2_consenso.csv
    # compat con el gold v1 (sin clase_dominante): deriva argmax y reporta empates
    python scripts/ingest_gold.py --books annotation/gold_set_v1.xlsx --allow-argmax
"""

import argparse
import csv
import json
from collections import Counter
from pathlib import Path

import openpyxl

from src.agents.gold.agreement import krippendorff_alpha, percent_agreement
from src.core.schema import CLASS_TO_IDX, IDEOLOGY_CLASSES

AXES = IDEOLOGY_CLASSES
DOMINANT_COL = "clase_dominante"


def read_book(path: Path) -> dict[str, dict]:
    """Lee la hoja 'Articulos' de un libro anotado.

    Returns:
        {id: {"dominant": str|None, "scores": {eje: int|None}, "notes": str}}
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "Articulos" not in wb.sheetnames:
        raise ValueError(f"{path}: no tiene hoja 'Articulos'")
    ws = wb["Articulos"]

    rows = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h is not None else "" for h in next(rows)]
    col = {name: i for i, name in enumerate(headers)}

    required = ["id", *AXES]
    missing = [c for c in required if c not in col]
    if missing:
        raise ValueError(f"{path}: faltan columnas {missing}")
    has_dominant = DOMINANT_COL in col

    annotations: dict[str, dict] = {}
    for row in rows:
        if row is None or col["id"] >= len(row) or not row[col["id"]]:
            continue
        article_id = str(row[col["id"]]).strip()

        scores: dict[str, int | None] = {}
        for axis in AXES:
            value = row[col[axis]] if col[axis] < len(row) else None
            if isinstance(value, (int, float)):
                scores[axis] = int(value)
            elif isinstance(value, str) and value.strip().isdigit():
                # Número con formato de texto (pegado que salta la data
                # validation de Excel): antes se descartaba en silencio.
                scores[axis] = int(value.strip())
            else:
                scores[axis] = None

        dominant = None
        if has_dominant and col[DOMINANT_COL] < len(row):
            raw = row[col[DOMINANT_COL]]
            if isinstance(raw, str) and raw.strip().lower() in CLASS_TO_IDX:
                dominant = raw.strip().lower()

        notes = ""
        if "notes" in col and col["notes"] < len(row) and row[col["notes"]]:
            notes = str(row[col["notes"]])

        annotations[article_id] = {
            "dominant": dominant,
            "scores": scores,
            "notes": notes,
        }
    wb.close()
    return annotations


def argmax_with_ties(scores: dict[str, int | None]) -> tuple[str | None, bool]:
    """(clase_argmax, hay_empate) desde escalas 1-5. None si no hay scores."""
    filled = {axis: v for axis, v in scores.items() if v is not None}
    if not filled:
        return None, False
    top = max(filled.values())
    winners = [axis for axis, v in filled.items() if v == top]
    return winners[0], len(winners) > 1


def resolve_dominant(
    annotation: dict, allow_argmax: bool,
) -> tuple[str | None, str | None]:
    """(clase, método): dominante explícita > argmax (solo si --allow-argmax)."""
    if annotation["dominant"]:
        return annotation["dominant"], "explicit"
    if allow_argmax:
        label, tie = argmax_with_ties(annotation["scores"])
        if label and not tie:
            return label, "argmax"
        if label and tie:
            return None, "tie"
    return None, None


def main() -> None:
    parser = argparse.ArgumentParser(description="Ingesta del gold anotado + Krippendorff α")
    parser.add_argument(
        "--books", nargs="+", required=True,
        help="Libros anotados (.xlsx), uno por anotador",
    )
    parser.add_argument(
        "--jsonl", type=str, default=None,
        help="JSONL con el contenido de los artículos del gold "
             "(default: se infiere del nombre del libro, ej. gold_set_v2.jsonl)",
    )
    parser.add_argument(
        "--consensus", type=str, default=None,
        help="CSV de consenso (columnas: id,label_final) para resolver discrepancias",
    )
    parser.add_argument(
        "--allow-argmax", action="store_true",
        help="Compat gold v1: deriva la dominante con argmax de las escalas 1-5 "
             "(los empates quedan pendientes de consenso). Preferir siempre v2.",
    )
    parser.add_argument(
        "--audit-silver", type=str, default=None,
        help="JSONL silver para auditar el juez LLM contra el consenso humano",
    )
    parser.add_argument("--output", type=str, default=None, help="JSONL de salida")
    args = parser.parse_args()

    from src.core.paths import ROOT

    book_paths = [Path(b) for b in args.books]
    for p in book_paths:
        if not p.exists():
            print(f"✗ No existe {p}")
            return

    # Inferir base: gold_set_v2_kevin.xlsx → gold_set_v2. El sufijo es el
    # NOMBRE DEL ANOTADOR y solo existe con varios libros; recortarlo con un
    # solo libro rompía el uso v1 (gold_set_v1 → gold_set).
    stem = book_paths[0].stem
    base = stem.rsplit("_", 1)[0] if len(book_paths) > 1 else stem
    annotation_dir = ROOT / "annotation"

    jsonl_path = Path(args.jsonl) if args.jsonl else annotation_dir / f"{base}.jsonl"
    output_path = Path(args.output) if args.output else annotation_dir / f"{base}_labeled.jsonl"
    discrepancies_path = annotation_dir / f"{base}_discrepancias.csv"
    report_path = annotation_dir / f"{base}_agreement_report.md"

    if not jsonl_path.exists():
        print(f"✗ No existe {jsonl_path} (contenido de los artículos). Usa --jsonl.")
        return
    articles: dict[str, dict] = {}
    with open(jsonl_path, encoding="utf-8") as f:
        for line in f:
            if line.strip():
                a = json.loads(line)
                articles[a["id"]] = a

    # --- Leer libros ---
    annotators = [p.stem.rsplit("_", 1)[-1] for p in book_paths]
    if len(set(annotators)) != len(annotators):
        print(f"✗ Nombres de anotador duplicados en los libros: {annotators}. "
              "Renombra los archivos (el sufijo tras el último '_' identifica "
              "al anotador).")
        return
    books = {name: read_book(p) for name, p in zip(annotators, book_paths)}
    for name, ann in books.items():
        n_dominant = sum(1 for a in ann.values() if a["dominant"])
        print(f"Libro {name}: {len(ann)} filas, {n_dominant} con {DOMINANT_COL}")

    # --- Concordancia sobre el solape ---
    all_ids = set().union(*(set(b) for b in books.values()))
    overlap_ids = [
        aid for aid in all_ids
        if sum(1 for b in books.values() if aid in b) >= 2
    ]

    report_lines = [
        "# Reporte de concordancia — gold set",
        "",
        f"- Libros: {', '.join(str(p) for p in book_paths)}",
        f"- Artículos anotados (unión): {len(all_ids)}",
        f"- Solape (2+ anotadores): {len(overlap_ids)}",
        "",
    ]

    alpha_dominant = None
    if overlap_ids:
        dominant_units = []
        for aid in overlap_ids:
            values = [
                resolve_dominant(b[aid], args.allow_argmax)[0]
                for b in books.values() if aid in b
            ]
            dominant_units.append(values)
        alpha_dominant = krippendorff_alpha(dominant_units, level="nominal")
        agreement = percent_agreement(dominant_units)

        report_lines += [
            "## Clase dominante (nominal)",
            "",
            f"- **Krippendorff α = {alpha_dominant:.4f}**" if alpha_dominant is not None
            else "- α no calculable (sin unidades pareables con dominante)",
            f"- Acuerdo bruto: {agreement:.1%}" if agreement is not None else "",
            f"- Umbral del anteproyecto: α ≥ 0.8 → "
            f"{'CUMPLE ✓' if (alpha_dominant or 0) >= 0.8 else 'NO CUMPLE ✗ (iterar codebook)'}",
            "",
            "## Escalas 1-5 por eje (interval)",
            "",
            "| Eje | α |",
            "|-----|---|",
        ]
        for axis in AXES:
            axis_units = [
                [b[aid]["scores"][axis] for b in books.values() if aid in b]
                for aid in overlap_ids
            ]
            alpha_axis = krippendorff_alpha(axis_units, level="interval")
            report_lines.append(
                f"| {axis} | {alpha_axis:.4f} |" if alpha_axis is not None
                else f"| {axis} | n/a |"
            )
        report_lines.append("")
    else:
        report_lines.append("⚠ Sin solape entre anotadores: α no calculable. "
                            "El anteproyecto exige doble anotación independiente "
                            "de una muestra compartida.")

    # --- Consenso previo (si existe) ---
    consensus: dict[str, str] = {}
    if args.consensus:
        with open(args.consensus, encoding="utf-8", newline="") as f:
            for row in csv.DictReader(f):
                label = (row.get("label_final") or "").strip().lower()
                if label in CLASS_TO_IDX:
                    consensus[row["id"].strip()] = label

    # --- Resolver etiqueta final por artículo ---
    final: dict[str, dict] = {}
    pending: list[dict] = []
    for aid in sorted(all_ids):
        per_annotator = {
            name: resolve_dominant(b[aid], args.allow_argmax)
            for name, b in books.items() if aid in b
        }
        labels = [label for label, _m in per_annotator.values() if label]
        distinct = set(labels)

        if aid in consensus:
            final[aid] = {"label": consensus[aid], "source": "human-consensus"}
        elif len(distinct) == 1 and labels:
            methods = {m for _l, m in per_annotator.values()}
            source = "human" if methods == {"explicit"} else "human-argmax"
            # "-agreed" exige que DOS anotadores hayan etiquetado de verdad:
            # con uno en blanco, resolver con la etiqueta del otro no es
            # acuerdo y marcarlo así inflaba la trazabilidad del α.
            if len(labels) >= 2:
                source += "-agreed"
            final[aid] = {"label": labels[0], "source": source}
        else:
            # Discrepancia o sin etiqueta → pendiente de consenso
            article = articles.get(aid, {})
            pending.append({
                "id": aid,
                "title": article.get("title", ""),
                "url": article.get("url", ""),
                **{
                    f"label_{name}": (label or method or "SIN_ETIQUETA")
                    for name, (label, method) in per_annotator.items()
                },
                "label_final": "",
            })

    # --- CSV de discrepancias / pendientes ---
    if pending:
        fieldnames = sorted({k for row in pending for k in row}, key=lambda k: (k != "id", k))
        with open(discrepancies_path, "w", encoding="utf-8", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(pending)
        print(f"\n⚠ {len(pending)} artículos requieren consenso → {discrepancies_path}")
        print("  Llenen `label_final` en sesión conjunta y re-corran con --consensus.")

    # --- JSONL etiquetado (solo artículos resueltos) ---
    written = 0
    label_dist: Counter = Counter()
    with open(output_path, "w", encoding="utf-8") as f:
        for aid, resolution in final.items():
            article = articles.get(aid)
            if article is None:
                print(f"  ⚠ {aid} anotado pero no está en {jsonl_path.name}; omitido")
                continue
            scores_by_annotator = {
                name: b[aid]["scores"] for name, b in books.items() if aid in b
            }
            record = {
                **article,
                "label": resolution["label"],
                "label_idx": CLASS_TO_IDX[resolution["label"]],
                "label_source": resolution["source"],
                "annotators": sorted(scores_by_annotator),
                "human_scores": scores_by_annotator,
            }
            f.write(json.dumps(record, ensure_ascii=False) + "\n")
            written += 1
            label_dist[resolution["label"]] += 1

    # --- Auditoría del silver (P1.3) ---
    if args.audit_silver and written:
        silver_path = Path(args.audit_silver)
        silver: dict[str, str] = {}
        with open(silver_path, encoding="utf-8") as f:
            for line in f:
                if not line.strip():
                    continue
                record = json.loads(line)
                label = record.get("label")
                if not label:
                    scores = {c: float(record.get(c, 0.0)) for c in AXES}
                    label = max(scores, key=scores.get)
                silver[record["id"]] = label

        audit_units, hits, total = [], 0, 0
        for aid, resolution in final.items():
            if aid in silver and aid in articles:
                audit_units.append([resolution["label"], silver[aid]])
                hits += int(resolution["label"] == silver[aid])
                total += 1
        if total:
            alpha_silver = krippendorff_alpha(audit_units, level="nominal")
            report_lines += [
                "## Auditoría del Silver Set (juez LLM vs consenso humano)",
                "",
                f"- Artículos comparados: {total}",
                f"- Accuracy del juez: {hits / total:.1%}",
                f"- **Krippendorff α (nominal) = {alpha_silver:.4f}**"
                if alpha_silver is not None else "- α no calculable",
                "",
            ]

    # --- Distribución final ---
    report_lines += [
        "## Distribución de clases del gold etiquetado",
        "",
        "| Clase | n |",
        "|-------|---|",
        *[f"| {cls} | {label_dist.get(cls, 0)} |" for cls in IDEOLOGY_CLASSES],
        "",
        f"Total etiquetados: {written} · Pendientes de consenso: {len(pending)}",
    ]
    report_path.write_text("\n".join(report_lines), encoding="utf-8")

    print(f"\n✓ Gold etiquetado: {output_path} ({written} artículos)")
    print(f"✓ Reporte de concordancia: {report_path}")
    if alpha_dominant is not None:
        status = "CUMPLE ✓" if alpha_dominant >= 0.8 else "NO CUMPLE ✗"
        print(f"  Krippendorff α (dominante) = {alpha_dominant:.4f} → {status} (umbral 0.8)")
    print("\nSiguiente paso: python scripts/prepare_splits.py "
          f"--gold-labeled {output_path}")


if __name__ == "__main__":
    main()
